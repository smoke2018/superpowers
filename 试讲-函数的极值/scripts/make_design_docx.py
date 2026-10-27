# -*- coding: utf-8 -*-
"""生成《函数的极值》教学设计（严格按学院模板结构）"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/workspace/试讲-函数的极值/教学设计-函数的极值.docx"

doc = Document()

# ---------- 全局 ----------
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(2.4)
sec.right_margin = Cm(2.4)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def set_font(run, name_cn="宋体", name_en="Times New Roman", size=11,
             bold=False, color=None):
    run.font.name = name_en
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(text, size=11, bold=False, cn="宋体", en="Times New Roman",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_first=0, space_before=0,
         space_after=0, line=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Pt(indent_first)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    r = p.add_run(text)
    set_font(r, cn, en, size, bold)
    return p


def heading(num_text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(num_text)
    set_font(r, "黑体", "Times New Roman", size, True)
    return p


def sub_heading(text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.4
    r = p.add_run(text)
    set_font(r, "黑体", "Times New Roman", size, True)
    return p


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def cell_text(cell, text, size=10.5, bold=False, cn="宋体",
              align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.35):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.line_spacing = line
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    for i, seg in enumerate(text.split("\n")):
        if i > 0:
            p = cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.line_spacing = line
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
        r = p.add_run(seg)
        set_font(r, cn, "Times New Roman", size, bold)


def set_col_widths(table, widths_cm):
    """真正固定列宽：设置 tblW、tblLayout=fixed、tblGrid 与每个 tcW"""
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr

    # 1) tblW
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tblPr.append(tblW)

    # 2) tblLayout = fixed
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    # 3) tblGrid
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 567)))
        grid.append(gc)
    tbl.insert(list(tbl).index(tblPr) + 1, grid)

    # 4) 每个单元格宽度（按合并跨度累计，跳过重复的合并单元）
    for row in table.rows:
        idx = 0
        seen = set()
        for cell in row.cells:
            tc_id = id(cell._tc)
            if tc_id in seen:
                continue
            seen.add(tc_id)
            span = cell.grid_span if hasattr(cell, "grid_span") else 1
            w = sum(widths_cm[idx: idx + span])
            if w:
                cell.width = Cm(w)
            idx += span

    # 5) 行不跨页拆分
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            el = OxmlElement("w:cantSplit")
            trPr.append(el)


def repeat_header(table):
    """表头跨页重复"""
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_cell_margins(table, left=0.05, right=0.05, top=0.02, bottom=0.02):
    """设置整表单元格内边距（cm），避免窄列文字被边距挤到换行"""
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for tag, v in (("w:top", top), ("w:left", left),
                   ("w:bottom", bottom), ("w:right", right)):
        e = OxmlElement(tag)
        e.set(qn("w:w"), str(int(v * 567)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


# ---------- 封面信息表 ----------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(10)
r = title.add_run("2024级数学专业师范生技能训练教学设计")
set_font(r, "黑体", "Times New Roman", 16, True)

info = doc.add_table(rows=3, cols=6)
info.style = "Table Grid"
info.alignment = WD_TABLE_ALIGNMENT.CENTER
rows_data = [
    ["姓名", "", "学段", "高中", "授课班级", ""],
    ["职称", "", "学科", "数学", "授课时间", ""],
]
for i, row in enumerate(rows_data):
    for j, val in enumerate(row):
        cell_text(info.rows[i].cells[j], val, size=10.5,
                  bold=(j % 2 == 0), align=WD_ALIGN_PARAGRAPH.CENTER)
# 课题行：先设列宽，再合并第 2~6 格
set_col_widths(info, [2.0, 3.0, 1.8, 2.6, 2.6, 4.2])
row3 = info.rows[2]
cell_text(row3.cells[0], "课题", size=10.5, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER)
merged = row3.cells[1]
merged = merged.merge(row3.cells[5])
cell_text(merged,
          "§6.2 函数的极值（北师大版选择性必修第二册 第二章 导数及其应用）",
          size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
merged.width = Cm(14.2)

doc.add_paragraph().paragraph_format.space_after = Pt(0)

# ---------- 1. 教材分析 ----------
heading("1.教材分析")
sub_heading("1.1 教材地位分析")
para("本节课选自北师大版（2019）高中数学选择性必修第二册第二章《导数及其应用》中"
     "§6 用导数研究函数的性质的第2课时——函数的极值。从知识序列看，学生此前已经学习了"
     "导数的概念、基本初等函数的导数公式、导数的运算法则以及导数与函数单调性的关系，"
     "本节课是在“用导数刻画增减性”的基础上，进一步研究增减转换处的局部特征；"
     "此后还将学习函数的最大（小）值以及导数在实际问题中的应用。因此，极值既是单调性研究的"
     "自然延续，又是最值研究的必要基础，在“局部—整体”的认识链条中起着承上启下的作用。",
     indent_first=22)
para("从育人价值看，极值概念中“附近”“局部”的限定，要求学生把图象直观上升为精确的数学表达，"
     "是培养数学抽象与逻辑推理素养的良好载体；“f′(x₀)=0 是否一定取极值”的辨析，"
     "则提供了体会必要条件与充分条件差异的典型素材。", indent_first=22)

sub_heading("1.2 教材内容分析")
para("本节教材的编排路径是：先观察具体函数图象中“由增变减”“由减变增”的分界点，"
     "抽象出极大值、极小值的概念；再借助导数符号的变化给出判定方法；最后通过例题示范"
     "求极值的完整过程，并用 f(x)=x³ 说明 f′(x₀)=0 只是极值存在的必要条件。"
     "教材把“形”的直观与“数”的刻画结合起来，体现了数形结合的思想方法。", indent_first=22)
para("教学重点：极值的概念，以及利用导数求函数极值的方法。", indent_first=22)
para("教学难点：极值概念中“局部性”的理解；判定中“两侧符号变号”这一条件的必要性。",
     indent_first=22)

# ---------- 2. 教学目标 ----------
heading("2.教学目标")
sub_heading("2.1 基础知识与基本技能目标")
para("① 理解极大值、极小值与极值点的概念，能结合图象说明极值的局部性；", indent_first=22)
para("② 掌握利用导数判定极值的方法，能按规范步骤求出函数的极大值与极小值；", indent_first=22)
para("③ 能辨析“f′(x₀)=0”与“x₀ 为极值点”的关系，并能举例说明二者的区别。", indent_first=22)

sub_heading("2.2 数学思想目标")
para("通过“图象特征—符号刻画”的相互转化，渗透数形结合与转化化归思想；"
     "通过“f′(x₀)=0 但 x₀ 不是极值点”的反例，渗透用反例检验命题、辨析条件的思维方法。",
     indent_first=22)

sub_heading("2.3 数学活动经验目标")
para("经历“观察图象—发现特征—抽象概念—建立判定—应用反思”的完整探究过程，"
     "积累从特殊到一般、从直观到严谨的数学活动经验，提升归纳概括与表达交流的能力。",
     indent_first=22)

sub_heading("2.4 核心素养目标")
para("① 数学抽象：从图象的“局部最高（低）点”抽象出极值概念，把握“局部”这一本质属性；",
     indent_first=22)
para("② 逻辑推理：由 f′(x) 在 x₀ 两侧的符号变化推出极值，辨析条件的充分性与必要性；",
     indent_first=22)
para("③ 数学运算：规范、准确地完成求导、解方程、判号与求值等运算。", indent_first=22)

# ---------- 3. 学情分析 ----------
heading("3.学情分析")
para("从认知基础看，学生已经掌握导数的概念与计算，会用 f′(x) 的符号判断函数的单调区间，"
     "这为本节课由“增减转换点”引出极值提供了直接支撑。从思维水平看，高二学生已具备一定的"
     "数形结合意识与归纳概括能力，能够从具体实例中提炼共同特征。", indent_first=22)
para("可能存在的困难有三点：一是容易把“极大值”误认为“函数的最大值”，忽视极值的局部性；"
     "二是容易形成“令 f′(x)=0 就得到极值点”的错误程序，忽略两侧符号的检验；"
     "三是对“极值点必须在定义域内部取得”这一限定缺乏认识。教学中应通过追问与反例逐一突破。",
     indent_first=22)

# ---------- 4. 教学策略分析 ----------
heading("4.教学策略分析")
para("遵循“以学生为主体，教师是数学课堂活动的组织者、引导者与合作者”的教学原则，"
     "采用启发式讲授与直观演示相结合的方法，以问题链贯穿全课：先用图象直观暴露“增减转换点”"
     "的特殊地位，再用导数符号给出精确判定，最后用反例完善认识。学生主要采用观察发现与"
     "自主探究的学习方式，在知识的“最近发展区”内完成概念的自主建构。教学中借助多媒体呈现"
     "函数图象以增强直观性，同时用板书呈现概念表述、判定方法与例题的完整解题过程，"
     "做到“技术辅助直观、板书承载严谨”。", indent_first=22)

# ---------- 5. 教学过程 ----------
p5_head = heading("5.教学过程")
p5_head.paragraph_format.keep_with_next = True

proc = doc.add_table(rows=1, cols=3)
proc.style = "Table Grid"
proc.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = proc.rows[0].cells
for i, t in enumerate(["教学内容与问题设计", "教师行为", "学生活动"]):
    cell_text(hdr[i], t, size=10.5, bold=True, cn="黑体",
              align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(hdr[i], "DCE6F1")


def add_teaching_row(content, teacher, student):
    row = proc.add_row()
    cell_text(row.cells[0], content, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    cell_text(row.cells[1], teacher, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    cell_text(row.cells[2], student, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    return row


def add_intent_row(text):
    row = proc.add_row()
    merged = row.cells[0].merge(row.cells[1]).merge(row.cells[2])
    cell_text(merged, text, size=10, cn="楷体", align=WD_ALIGN_PARAGRAPH.LEFT)
    shade(merged, "F2F2F2")


add_teaching_row(
    "5.1 创设情境（约3分钟）\n"
    "展示函数 f(x)=x³-3x 的图象。\n"
    "问题1：你能求出 f′(x)，并说出它的单调区间吗？\n"
    "问题2：图象在 x=-1 和 x=1 处分别发生了什么变化？\n"
    "问题3：像这样“由增变减”“由减变增”的分界点，"
    "其函数值在这一点“附近”处于什么地位？",
    "展示图象，依次抛出问题链；"
    "在学生回答后追问“是附近最大，还是整个定义域上最大”，"
    "引导学生用“附近”“局部”描述特征，并板书课题。",
    "观察图象，回忆用导数判断单调性的方法，"
    "回答 f′(x)=3(x+1)(x-1) 及各单调区间；"
    "尝试用语言描述两个分界点的特征。")
add_intent_row(
    "设计意图：以学生已掌握的单调性为起点，用图象直观制造“增减转换点究竟特殊在哪里”的认知需求，"
    "使极值概念的产生具有逻辑必然性；在描述中反复使用“附近”一词，为后续概念的精确表述作语言铺垫，"
    "同时渗透数形结合思想。")

add_teaching_row(
    "5.2 新知探究（约6分钟）\n"
    "（1）极值的概念\n"
    "设函数 y=f(x) 在点 x₀ 附近有定义。若对 x₀ 附近的所有点 x，都有 f(x)<f(x₀)，"
    "则称 f(x₀) 为函数的一个极大值，x₀ 为极大值点；若都有 f(x)>f(x₀)，"
    "则称 f(x₀) 为极小值，x₀ 为极小值点。极大值与极小值统称极值。\n"
    "关键词辨析：①“附近”——极值是局部性质；②极值点是自变量的值，极值是函数值；"
    "③极值点须在定义域内部取得。\n"
    "（2）极值的判定\n"
    "若 f′(x₀)=0：x₀ 左侧 f′(x)>0、右侧 f′(x)<0，则 f(x₀) 为极大值；"
    "左负右正，则为极小值；两侧同号，则 f(x₀) 不是极值。\n"
    "（3）辨析与追问\n"
    "问题4：若 f′(x₀)=0，x₀ 一定是极值点吗？请举例说明。\n"
    "问题5：极大值一定是函数的最大值吗？",
    "板书极值的定义与判定，逐个解释三个关键词；"
    "问题4出示 y=x³ 的图象，提示“导数在 x=0 两侧符号如何”；"
    "问题5回扣情境中的图象，引导观察 x→+∞ 时的函数值趋势。"
    "明确：f′(x₀)=0 是极值存在的必要条件，不是充分条件。",
    "对照定义复述“附近”的含义；"
    "举例：f(x)=x³，f′(0)=0，但 f(x) 在 R 上单调递增，"
    "故 x=0 不是极值点；"
    "回答：极大值 2 只是在 x=-1 附近最大，函数无最大值。")
add_intent_row(
    "设计意图：概念在图象直观中生成，判定在符号分析中建立，认识在反例辨析中完善。"
    "问题4直指学生的典型错误程序，用反例说明“必要非充分”，培养思维的严谨性；"
    "问题5区分极值与最值，防止概念混淆。三个关键词的逐一辨析，落实概念的内涵与外延。")

add_teaching_row(
    "5.3 例题讲解（约5分钟）\n"
    "例1 求函数 f(x)=x³-3x 的极值。\n"
    "① 分析题意：定义域为 R；求极值即找 f′(x)=0 的点，并判断其两侧 f′(x) 的符号。\n"
    "② 拟定计划：求导 → 解方程 f′(x)=0 → 列表判号 → 按判定写出极值。\n"
    "③ 实施计划：\n"
    "f′(x)=3x²-3=3(x+1)(x-1)，令 f′(x)=0，解得 x₁=−1，x₂=1。\n"
    "列表（见板书）：f′(x) 在 (-∞,-1) 上为正，(-1,1) 上为负，(1,+∞) 上为正。\n"
    "故 f(-1)=2 为极大值，f(1)=-2 为极小值。\n"
    "④ 回顾反思：极大值 2 不是最大值；符号变化与图象“先增后减再增”一致；"
    "若省略列表、只令 f′(x)=0，就可能把 y=x³ 中的 x=0 误判为极值点。",
    "按“分析—计划—实施—回顾”四步示范，"
    "完整板书解题过程与符号表；"
    "在“回顾”环节追问“为什么必须列表”，"
    "并请学生说明极大值与最大值的区别。",
    "跟随分析题意，独立写出求导与解方程过程；"
    "回答各区间上 f′(x) 的符号；"
    "回答“列表是为了检验两侧是否变号”。")
add_intent_row(
    "设计意图：以四步示范使解题过程结构化、可迁移，符合试讲对例题讲解“分析题意、拟定计划、"
    "实施计划、回顾反思”的要求；用表格把符号分析显性化，突出判定的核心是“两侧变号”；"
    "回顾环节回扣反例，使概念、判定与应用形成闭环。")

add_teaching_row(
    "5.4 课堂小结（约1分钟）\n"
    "以框图梳理研究路径：\n"
    "图象特征（增减转换）→ 极值概念（局部最大/最小）→ 判定方法（f′(x₀)=0 且两侧变号）"
    "→ 应用（求极值四步）。\n"
    "强调两点：极值是局部概念；f′(x₀)=0 未必是极值点。",
    "提出问题“这节课我们怎样认识了一个新概念”，"
    "与学生共同完成框图，板书关键节点。",
    "回顾学习过程，尝试用一句话概括研究路径；"
    "补充说明易错点。")
add_intent_row(
    "设计意图：用框图呈现“直观—概念—判定—应用”的研究路径，帮助学生形成结构化的认识，"
    "并把本节课的方法（由形到数、由特殊到一般、用反例检验）提升为可迁移的数学活动经验。")

add_teaching_row(
    "5.5 作业设计\n"
    "① 教材 §6.2 练习第1、2题；\n"
    "② 求函数 f(x)=x³-3x² 的极值，仿照例1写出完整过程；\n"
    "③ 思考：若 f′(x₀)=0 且在 x₀ 两侧 f′(x) 同号，函数图象在 x₀ 处有什么特点？",
    "布置作业并说明第②题的要求（步骤完整、表格规范）；"
    "第③题不作统一要求，供学有余力的学生思考。",
    "记录作业，明确书写规范。")
add_intent_row(
    "设计意图：作业分基础、巩固、拓展三个层次，兼顾不同水平的学生；第③题指向后续内容，"
    "为学习函数的图象特征埋下伏笔。")

set_col_widths(proc, [8.2, 3.9, 4.1])
repeat_header(proc)

# ---------- 6. 板书设计 ----------
heading("6.板书设计")
para("（黑板分三个区域，左侧概念与判定，中间例题，右侧图象与反例）", size=10,
     cn="楷体", indent_first=22)
board = doc.add_table(rows=1, cols=3)
board.style = "Table Grid"
board.alignment = WD_TABLE_ALIGNMENT.CENTER
b = board.rows[0].cells
cell_text(b[0],
          "函数的极值\n\n"
          "一、概念\n"
          "设 y=f(x) 在 x₀ 附近有定义：\n"
          "f(x)<f(x₀) → 极大值 f(x₀)\n"
          "f(x)>f(x₀) → 极小值 f(x₀)\n"
          "（局部、内部点）\n\n"
          "二、判定（f′(x₀)=0）\n"
          "左正右负 → 极大值\n"
          "左负右正 → 极小值\n"
          "两侧同号 → 不是极值\n\n"
          "三、注意\n"
          "f′(x₀)=0 ⇏ 极值点\n"
          "反例：y=x³，x=0", size=10, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.25)

# 中间：例题（题干 + 解答），符号表用嵌套表格保证对齐
cell_text(b[1], "例1 求 f(x)=x³-3x 的极值", size=10, bold=True,
          align=WD_ALIGN_PARAGRAPH.LEFT, line=1.25)
p = b[1].add_paragraph()
p.paragraph_format.line_spacing = 1.25
p.paragraph_format.space_before = Pt(2)
r = p.add_run("解：定义域 R\nf′(x)=3x²-3=3(x+1)(x-1)\n令 f′(x)=0，得 x₁=-1，x₂=1")
set_font(r, "宋体", "Times New Roman", 10)

sign = b[1].add_table(rows=3, cols=6)
sign.style = "Table Grid"
sign_data = [
    ["x", "(-∞,-1)", "-1", "(-1,1)", "1", "(1,+∞)"],
    ["f′(x)", "+", "0", "-", "0", "+"],
    ["f(x)", "↗", "极大", "↘", "极小", "↗"],
]
for i, rowvals in enumerate(sign_data):
    for j, val in enumerate(rowvals):
        cell_text(sign.rows[i].cells[j], val, size=7,
                  align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
set_cell_margins(sign, left=0.02, right=0.02, top=0.01, bottom=0.01)
set_col_widths(sign, [1.3, 1.15, 0.95, 1.1, 0.95, 1.25])

p = b[1].add_paragraph()
p.paragraph_format.line_spacing = 1.25
p.paragraph_format.space_before = Pt(3)
r = p.add_run("∴ 极大值 f(-1)=2\n　 极小值 f(1)=-2")
set_font(r, "宋体", "Times New Roman", 10)

cell_text(b[2],
          "（图象区）\n\n"
          "1. 作 y=x³-3x 图象，\n"
          "   标出 (-1,2)、(1,-2)\n"
          "   及水平切线\n\n"
          "2. 作 y=x³ 简图，\n"
          "   标出 f′(0)=0\n"
          "   但无极值", size=10, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.25)
set_col_widths(board, [4.6, 7.2, 4.4])

# ---------- 7. 教学反思 ----------
doc.add_page_break()
heading("7.教学反思")
sub_heading("三个亮点")
para("① 概念引入有根：以学生已掌握的单调性为起点，从图象的“增减转换点”抽象出极值，"
     "使新概念的生长具有逻辑必然性，避免了概念的“空降”。", indent_first=22)
para("② 反例使用有力：用 f(x)=x³ 突破“f′(x)=0 即极值点”的思维定式，"
     "把必要条件与充分条件的辨析落到实处，学生印象深刻。", indent_first=22)
para("③ 解题示范有序：例题按“分析—计划—实施—回顾”四步展开，符号表把“两侧变号”"
     "这一关键条件显性化，解题过程结构化、可迁移。", indent_first=22)

sub_heading("两个培育")
para("培育数学抽象素养：从图象的直观特征到概念的精确界定，经历“去情境、抓本质”的抽象过程。",
     indent_first=22)
para("培育逻辑推理素养：通过符号表与反例，体会结论成立的条件与边界，形成言必有据的习惯。",
     indent_first=22)

sub_heading("一个思考")
para("受课堂时间限制，“极值与最值”的区别仅作点明而未展开，可在后续最值课中安排对比；"
     "另外，学生对“附近”的严格含义（邻域）尚不能完全把握，如何用更贴近高中生语言的方式"
     "表达这一限定，仍需进一步打磨。若再上一次，可考虑把问题5调整到小结环节，"
     "以保证例题讲解的时间更充裕。", indent_first=22)

doc.save(OUT)
print("saved:", OUT)
