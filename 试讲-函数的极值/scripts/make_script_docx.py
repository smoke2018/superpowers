# -*- coding: utf-8 -*-
"""生成《函数的极值》试讲逐字稿（15分钟）"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/workspace/试讲-函数的极值/逐字稿-函数的极值.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.6)
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.4)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.paragraph_format.line_spacing = 1.45
style.paragraph_format.space_after = Pt(0)


def set_font(run, name_cn="宋体", name_en="Times New Roman", size=12,
             bold=False, color=None):
    run.font.name = name_en
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name_cn)
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)


def para(runs, size=12, line=1.45, before=0, after=0, indent_char=0,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent_char:
        pPr = p._p.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:firstLineChars"), str(int(indent_char * 100)))
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, opt in runs:
        r = p.add_run(text)
        set_font(r, opt.get("cn", "宋体"), opt.get("en", "Times New Roman"),
                 opt.get("size", size), opt.get("bold", False),
                 opt.get("color"))
    return p


def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, "黑体", "Times New Roman", 18, True, (0x1F, 0x38, 0x64))
    return p


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, "黑体", "Times New Roman", 13.5, True, (0x1F, 0x38, 0x64))
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, "黑体", "Times New Roman", 11.5, True)
    return p


def say(text):
    """教师讲授语"""
    return para([("【师】", {"bold": True, "cn": "黑体"}),
                 (text, {})], indent_char=2)


def stu(text):
    """预设学生回答"""
    p = para([("【生】", {"bold": True, "cn": "黑体", "color": (0x2E, 0x53, 0x95)}),
              (text, {"color": (0x2E, 0x53, 0x95)})], indent_char=2)
    p.paragraph_format.keep_with_next = True
    return p


def act(text):
    """动作、板书、PPT 提示"""
    p = para([("〔" + text + "〕", {"cn": "楷体", "size": 10.5,
                                   "color": (0x6B, 0x73, 0x80)})],
             indent_char=2, before=2, after=2)
    p.paragraph_format.keep_with_next = True
    return p


# ================= 标题与说明 =================
title("《函数的极值》试讲逐字稿")
para([("北师大版（2019）选择性必修第二册　第二章 §6.2　总时长 15 分钟",
       {"size": 10.5, "color": (0x6B, 0x73, 0x80),
        "align": WD_ALIGN_PARAGRAPH.CENTER})],
     align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

para([("使用说明：", {"bold": True, "cn": "黑体", "size": 10.5}),
      ("【师】为教师讲授语，建议按原话讲，概念表述不可改动；【生】为预设学生回答，"
       "学生答出即接下一句，答不出按括号内的追问处理；〔  〕为板书、课件与动作提示，"
       "不念出来。全稿正文约 2500 字，按每分钟 200 字计约 12 分钟，"
       "余下时间留给板书书写、学生思考与练习。", {"size": 10.5})],
     indent_char=0, after=10)

# ================= 时间轴 =================
h1("时间分配一览")
rows = [
    ("环节", "起止", "时长", "任务"),
    ("一、导入", "0:00—2:40", "2分40秒", "复习单调性 → 观察图象 → 发现“增减转换点”"),
    ("二、新知探究", "2:40—8:40", "6分钟", "抽象概念 → 辨析关键词 → 判定方法 → 反例"),
    ("三、例题讲解", "8:40—13:20", "4分40秒", "分析题意 → 拟定计划 → 实施计划 → 回顾反思"),
    ("四、小结与作业", "13:20—15:00", "1分40秒", "结构化小结 → 两个提醒 → 分层作业"),
]
tbl = doc.add_table(rows=len(rows), cols=4)
tbl.style = "Table Grid"
widths = [3.2, 3.4, 2.6, 6.6]
for j, w in enumerate(widths):
    tbl.columns[j].width = Cm(w)
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.width = Cm(widths[j])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 3 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(val)
        set_font(r, "宋体", "Times New Roman", 10, i == 0)
        if i == 0:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "E8EDF6")
            cell._tc.get_or_add_tcPr().append(shd)

# ================= 一、导入 =================
h1("一、导入（0:00—2:40，约 2 分 40 秒）")
act("课件第 3 页；黑板左侧留白，暂不写课题")
say("上课！同学们好，请坐。")
say("前面几节课，我们学习了用导数判断函数的单调性：在某个区间内，如果 f′(x)>0，"
    "函数在这个区间上单调递增；如果 f′(x)<0，函数在这个区间上单调递减。"
    "有了导数，我们不必画出完整的图象，也能看清函数的“升降”。")
say("请看大屏幕。我们来算一算函数 f(x)=x³−3x。求导，f′(x)=3x²−3=3(x+1)(x−1)。"
    "令 f′(x)>0，解得 x<−1 或 x>1，所以函数在 (−∞,−1) 和 (1,+∞) 上单调递增；"
    "令 f′(x)<0，解得 −1<x<1，所以函数在 (−1,1) 上单调递减。")
act("板书：f′(x) 的符号与单调区间")
say("现在请大家看图象。请一位同学回答：在 x=−1 和 x=1 这两处，图象发生了什么变化？")
stu("在 x=−1 处，函数由增变减；在 x=1 处，函数由减变增。")
say("说得很好。这两点，是函数“增减的转换点”。请大家再仔细看：这两个转换点附近的函数值，"
    "处在什么地位？")
stu("x=−1 附近的函数值，以这一点的函数值为最大；x=1 附近的函数值，以这一点的函数值为最小。")
say("“附近”“最大”“最小”——这三个词抓住了关键。在一点的附近取到最大或最小的函数值，"
    "数学上给它一个专门的名称。这就是我们今天要研究的内容——函数的极值。")
act("板书课题：§6.2 函数的极值")

# ================= 二、新知探究 =================
h1("二、新知探究（2:40—8:40，约 6 分钟）")

h2("（一）从图象到概念（2:40—4:30）")
act("课件第 5 页")
say("我们先把刚才的直观描述，翻译成准确的数学语言。")
say("设函数 y=f(x) 在点 x₀ 附近有定义。如果对 x₀ 附近的所有 x，都有 f(x)<f(x₀)，"
    "那么称 f(x₀) 为函数的一个极大值，x₀ 叫做极大值点；如果对 x₀ 附近的所有 x，"
    "都有 f(x)>f(x₀)，那么称 f(x₀) 为极小值，x₀ 叫做极小值点。"
    "极大值与极小值统称极值，极大值点与极小值点统称极值点。")
act("板书：定义，左右两栏对照写")
say("请大家把定义读一遍，边读边圈出两个词——“附近”和“所有”。"
    "谁来说说，这两个词为什么都不能省？")
stu("“附近”说明只和这一点周围比；“所有”说明周围的每一个点都要满足。")
say("完全正确。这两句话，说的正是极值的“局部性”。")

h2("（二）辨析三个关键词（4:30—6:00）")
act("课件第 6 页")
say("围绕定义，有三处最容易出错，我们逐一辨析。")
say("第一处，“附近”。极值是局部性质：它只与 x₀ 附近的函数值比较，不与整个定义域上的"
    "函数值比较。所以，极大值不一定是最大值。")
say("看这个例子：f(x)=x³−3x 在 x=−1 处取得极大值 2。可是当 x 趋向正无穷时，f(x) 也趋向"
    "正无穷，这个函数根本没有最大值。可见“极大”两个字，讲的是“附近最高”，不是“全局最高”。")
say("第二处，“极值点”与“极值”。极值点 x₀ 是自变量的取值，极值 f(x₀) 是函数值，"
    "这是两个不同的对象。回答问题时要说清楚：极大值点是 x=−1，极大值是 2。")
say("第三处，“内部点”。极值点必须在定义域内部取得，区间端点处不谈极值。"
    "因为端点只有一侧，谈不上“附近两侧”的比较。")
act("板书：极值点 ≠ 极值；端点不算极值点")
say("把这三个关键词连起来，极值概念的轮廓就清楚了。")

h2("（三）极值的判定（6:00—7:10）")
act("课件第 7 页")
say("定义是“看图说话”，可解题时往往没有现成的图象。能不能像判断单调性那样，"
    "用导数来判断极值呢？")
say("回到刚才的图象。在 x=−1 的左侧，函数递增，f′(x)>0；在右侧，函数递减，f′(x)<0。"
    "导数由正变负，函数就在这里“由升转降”，形成一个“山顶”，取得极大值。")
say("同样，在 x=1 处导数由负变正，函数“由降转升”，形成一个“谷底”，取得极小值。")
say("于是得到判定方法：设 f′(x₀)=0。如果 x₀ 左侧 f′(x)>0、右侧 f′(x)<0，则 f(x₀) 是极大值；"
    "如果左侧 f′(x)<0、右侧 f′(x)>0，则 f(x₀) 是极小值；如果两侧符号相同，则 f(x₀) 不是极值。")
act("板书：判定表；口诀“左正右负取极大，左负右正取极小；符号不变，极值不现”")

h2("（四）反例辨析：必要还是充分（7:10—8:40）")
act("课件第 8 页")
say("听到这里，可能有同学会想：那我只要令 f′(x)=0，解出来的点就都是极值点，对吗？")
say("我们看一个反例：f(x)=x³。f′(x)=3x²，f′(0)=0。但是当 x<0 时 f′(x)>0，当 x>0 时"
    "f′(x)>0，两侧同号。函数在 R 上始终单调递增，图象只是一路“爬升”，"
    "在 x=0 处短暂变平，并没有形成山顶或谷底。所以 x=0 不是极值点。")
act("追问：x=0 左右两侧的导数各是什么符号？")
stu("左边是正的，右边也是正的，两侧同号。")
say("两侧同号，说明函数的增减方向在 x=0 并没有被“逆转”。所以 f′(x₀)=0 只是极值存在的必要条件，"
    "不是充分条件。求极值时，只解方程还不够，必须检验两侧的符号是否变号。")
act("板书：f′(x₀)=0 ⇏ 极值点；反例 y=x³，x=0")

# ================= 三、例题讲解 =================
h1("三、例题讲解（8:40—13:20，约 4 分 40 秒）")
act("课件第 9 页")
say("下面我们用这套方法解决一个具体问题。例 1：求函数 f(x)=x³−3x 的极值。")
say("第一步，分析题意。函数的定义域是 R。求极值，就是先找使 f′(x)=0 的点，"
    "再判断每个点两侧 f′(x) 的符号。这里要注意：题目没有给图象，我们要靠导数“算出”"
    "函数的升降。")
say("第二步，拟定计划。求导，解方程 f′(x)=0，列表判断符号，再按判定写出极值。"
    "四步走，一步不落。")
say("第三步，实施计划。")
act("板书解题过程；课件第 10 页")
say("求导：f′(x)=3x²−3=3(x+1)(x−1)。令 f′(x)=0，解得 x₁=−1，x₂=1。")
say("这两个点把定义域分成三个区间：(−∞,−1)、(−1,1)、(1,+∞)。列表：在 (−∞,−1) 上，"
    "f′(x)>0；在 (−1,1) 上，f′(x)<0；在 (1,+∞) 上，f′(x)>0。")
act("板书符号表：先写 x 的范围，再逐行填 f′(x) 的符号，最后写 f(x) 的增减")
say("所以在 x=−1 处，导数由正变负，函数取得极大值 f(−1)=(−1)³−3×(−1)=2；"
    "在 x=1 处，导数由负变正，函数取得极小值 f(1)=1³−3×1=−2。"
    "所以，f(x) 的极大值为 2，极小值为 −2。")
say("第四步，回顾反思。这里有三个问题。")
act("课件第 11 页")
say("问题一：极大值 2 是最大值吗？不是。x 趋向正无穷时 f(x) 趋向正无穷，函数没有最大值。"
    "这正好印证了极值的局部性。")
say("问题二：这个结果可信吗？符号表显示函数“先增后减再增”，与我们画的图象完全一致，"
    "数与形互相印证。")
say("问题三：表格能省略吗？不能。如果只令 f′(x)=0 就下结论，就可能把 y=x³ 中的 x=0 "
    "误判为极值点。表格的作用，就是把“两侧是否变号”这个条件摆到明面上。")
say("请同学们在练习本上，按这四步把本题完整地写一遍。")
act("留白约 30 秒，巡视，提醒书写规范：定义域、求导、解方程、列表、结论")

# ================= 四、小结与作业 =================
h1("四、课堂小结与作业（13:20—15:00，约 1 分 40 秒）")
act("课件第 12 页")
say("这节课我们学习了什么？我们从图象上看，极值出现在“增减转换点”；"
    "从概念上看，极值是“局部最大或局部最小”；从方法上看，判定极值要“f′(x₀)=0 且两侧变号”；"
    "从应用上看，求极值有四步——分析题意、拟定计划、实施计划、回顾反思。")
say("两个提醒：第一，极值是局部概念，极大值不一定是最大值；"
    "第二，f′(x₀)=0 未必是极值点，必须检验两侧的符号。")
say("我们研究问题的路径是：由形到数、由特殊到一般、用反例检验结论的边界。"
    "这也是数学研究的一般方法。")
act("课件第 13 页")
say("作业：第一，教材 §6.2 练习第 1、2 题；第二，求函数 f(x)=x³−3x² 的极值，"
    "仿照例 1 写出完整过程；第三，思考：若 f′(x₀)=0，且在 x₀ 两侧 f′(x) 同号，"
    "函数图象在 x₀ 处有什么特点？")
say("下节课，我们从“局部”走向“整体”，研究函数的最大值与最小值。下课！")

# ================= 附：节奏提示 =================
h1("附：节奏与临场提示")
notes = [
    ("导入（0:00—2:40）", "语速稍慢；“附近”“最大”“最小”三个词要重读；"
     "在最后一个问题之前，绝不要提前说出“极值”二字，把发现的机会留给学生。"),
    ("新知（2:40—8:40）", "概念表述必须一字不差，宁可慢，不可错；三个关键词各用两三句话点透，"
     "不展开举例；反例处要有停顿，等学生产生疑问再给反例。"),
    ("例题（8:40—13:20）", "边讲边写，符号表分三行写，写完再讲结论；"
     "“回顾反思”三问按顺序问，每问不超过 20 秒，不要在此处展开新内容。"),
    ("小结（13:20—15:00）", "只讲结构与两个提醒，不重复推导；作业读完即可，不再解释。"),
    ("超时预案", "删去“第二处、第三处”关键词中的举例，或把回顾反思的第 2 问并入小结，"
     "可压缩约 40 秒。"),
]
for name, txt in notes:
    para([("◆ " + name + "：", {"bold": True, "cn": "黑体", "size": 10.5}),
          (txt, {"size": 10.5})], indent_char=0, line=1.4, after=3)

doc.save(OUT)
print("saved:", OUT)
